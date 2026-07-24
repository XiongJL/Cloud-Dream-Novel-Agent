from __future__ import annotations

from pathlib import Path
import json
import subprocess
import sys

import pytest

from document_extractor.extractor import DocumentExtractorError, extract_document


def request(path: Path, *, max_characters: int = 2_000_000) -> dict[str, object]:
    return {
        "protocolVersion": "document-extractor-v1",
        "jobId": "job-test",
        "temporaryFilePath": str(path),
        "originalFileName": path.name,
        "extension": path.suffix,
        "limits": {"maxCharacters": max_characters, "timeoutMs": 30_000},
    }


def test_extracts_markdown_with_stable_offsets(tmp_path: Path) -> None:
    source = tmp_path / "outline.md"
    source.write_text("# Main line\n\nOpening paragraph.\n\n## Conflict\n\n- First beat", encoding="utf-8")

    first = extract_document(request(source))
    second = extract_document(request(source))

    assert first == second
    assert first["title"] == "Main line"
    assert [block["type"] for block in first["blocks"]] == ["heading", "paragraph", "heading", "list_item"]
    for block in first["blocks"]:
        assert first["plainText"][block["startOffset"]:block["endOffset"]] == block["text"]


def test_reports_truncation(tmp_path: Path) -> None:
    source = tmp_path / "long.txt"
    source.write_text("abcdefghij", encoding="utf-8")
    result = extract_document(request(source, max_characters=5))
    assert result["plainText"] == "abcde"
    assert result["metadata"]["warnings"][0]["code"] == "CONTENT_TRUNCATED"


def test_offsets_use_javascript_utf16_code_units(tmp_path: Path) -> None:
    source = tmp_path / "emoji.md"
    source.write_text("# 角色😀\n\n表情之后的内容", encoding="utf-8")

    result = extract_document(request(source))

    assert result["blocks"][0]["text"] == "角色😀"
    assert result["blocks"][0]["startOffset"] == 0
    assert result["blocks"][0]["endOffset"] == 4
    assert result["blocks"][1]["startOffset"] == 6


def test_rejects_unsupported_extension(tmp_path: Path) -> None:
    source = tmp_path / "notes.rtf"
    source.write_text("content", encoding="utf-8")
    with pytest.raises(DocumentExtractorError, match="Unsupported") as captured:
        extract_document(request(source))
    assert captured.value.code == "UNSUPPORTED_FILE_TYPE"


def test_extracts_docx_structure(tmp_path: Path) -> None:
    from docx import Document

    source = tmp_path / "characters.docx"
    document = Document()
    document.add_heading("Characters", level=1)
    document.add_paragraph("Lin is the protagonist.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Role"
    document.save(source)

    result = extract_document(request(source))
    assert [block["type"] for block in result["blocks"]] == ["heading", "paragraph", "table"]
    assert "Name\tRole" in result["plainText"]


def test_extracts_pdf_pages(tmp_path: Path) -> None:
    import fitz

    source = tmp_path / "research.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Evidence on page one")
    document.save(source)
    document.close()

    result = extract_document(request(source))
    assert result["metadata"]["pageCount"] == 1
    assert result["blocks"][0]["page"] == 1
    assert "Evidence on page one" in result["plainText"]


def test_cli_returns_versioned_envelope(tmp_path: Path) -> None:
    source = tmp_path / "notes.txt"
    source.write_text("Readable notes 😀", encoding="utf-8")
    completed = subprocess.run(
        [sys.executable, "-m", "document_extractor"],
        input=json.dumps(request(source)).encode("utf-8"),
        capture_output=True,
        check=True,
        cwd=Path(__file__).resolve().parents[1],
    )
    envelope = json.loads(completed.stdout.decode("utf-8"))
    assert envelope["protocolVersion"] == "document-extractor-v1"
    assert envelope["jobId"] == "job-test"
    assert envelope["ok"] is True
    assert envelope["document"]["plainText"] == "Readable notes 😀"
