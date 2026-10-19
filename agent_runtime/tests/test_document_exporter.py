from __future__ import annotations

from copy import deepcopy
import json
import os
from pathlib import Path
import subprocess
import sys
from time import perf_counter
from zipfile import ZipFile

from docx import Document
from lxml import etree
import pytest

from document_exporter.exporter import PROTOCOL, export_document


def book() -> dict:
    return {
        "title": "长夜😀", "author": "行舟", "description": "首行简介\n\n第二段简介",
        "volumes": [
            {"id": "v1", "title": "第 1 卷 江湖", "chapters": [
                {"id": "c1", "title": "第 1 章 重名", "paragraphs": ["　　正文😀  ", "", "段内\n换行\t尾声"]},
                {"id": "c2", "title": "第 2 章 重名", "paragraphs": []},
            ]},
            {"id": "v2", "title": "第 2 卷 空卷", "chapters": []},
        ],
    }


def test_txt_metadata_order_bom_and_whitespace(tmp_path: Path) -> None:
    output = tmp_path / "book.txt"
    progress = []
    assert export_document(book(), "txt", output, progress.append) == {"volumes": 2, "chapters": 2, "empty": False}
    raw = output.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    text = raw.decode("utf-8-sig")
    assert text.startswith("长夜😀\r\n\r\n作者：行舟\r\n\r\n作品简介\r\n\r\n首行简介\r\n\r\n第二段简介")
    assert "　　正文😀  \r\n\r\n段内\r\n换行\t尾声" in text
    assert "暂无" not in text
    assert text.endswith("第 2 卷 空卷\r\n")
    assert b"\n" not in raw.replace(b"\r\n", b"")
    assert progress == [1, 2]


def test_docx_headings_cached_links_toc_and_body_match(tmp_path: Path) -> None:
    source = book()
    output = tmp_path / "book.docx"
    export_document(source, "docx", output)
    doc = Document(output)
    assert doc.core_properties.author == "行舟"
    paragraphs = [p.text for p in doc.paragraphs]
    assert paragraphs[:6] == ["长夜😀", "作者：行舟", "作品简介", "首行简介", "", "第二段简介"]
    assert [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")] == [
        ("Heading 1", "第 1 卷 江湖"), ("Heading 2", "第 1 章 重名"), ("Heading 2", "第 2 章 重名"), ("Heading 1", "第 2 卷 空卷")
    ]
    assert "　　正文😀  " in paragraphs
    assert "段内\n换行\t尾声" in paragraphs
    with ZipFile(output) as archive:
        xml = etree.fromstring(archive.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        anchors = xml.xpath("//w:hyperlink/@w:anchor", namespaces=ns)
        bookmarks = xml.xpath("//w:bookmarkStart/@w:name", namespaces=ns)
        assert len(set(anchors)) == len(anchors) == 4
        assert set(anchors) == set(bookmarks)
        assert 'TOC \\o "1-2" \\h' in "".join(xml.xpath("//w:instrText/text()", namespaces=ns))
        assert xml.xpath("//w:fldChar/@w:fldCharType", namespaces=ns) == ["begin", "separate", "end"]
        assert b"updateFields" in archive.read("word/settings.xml")
        for style in ("Heading1", "Heading2"):
            styles = etree.fromstring(archive.read("word/styles.xml"))
            assert styles.xpath(f'//w:style[@w:styleId="{style}"]//w:outlineLvl', namespaces=ns)


@pytest.mark.parametrize("format", ["txt", "docx"])
def test_empty_metadata_and_no_chapters(tmp_path: Path, format: str) -> None:
    snapshot = {"title": "空作品", "author": "", "description": "", "volumes": []}
    output = tmp_path / f"book.{format}"
    assert export_document(snapshot, format, output)["empty"]
    if format == "txt":
        assert output.read_bytes().decode("utf-8-sig") == "空作品\r\n"
    else:
        assert [p.text for p in Document(output).paragraphs] == ["空作品"]


def test_worker_protocol_and_bad_request(tmp_path: Path) -> None:
    (tmp_path / "snapshot.json").write_text(json.dumps(book()), encoding="utf-8")
    request = {"protocolVersion": PROTOCOL, "jobId": "job-😀", "directory": str(tmp_path), "format": "docx"}
    result = subprocess.run([sys.executable, "-m", "document_exporter"], input=json.dumps(request).encode(), capture_output=True, timeout=30, cwd=Path(__file__).resolve().parents[1])
    assert result.returncode == 0, result.stderr
    events = [json.loads(line) for line in result.stdout.decode().splitlines()]
    assert all(e["jobId"] == "job-😀" and e["protocolVersion"] == PROTOCOL for e in events)
    assert events[-1]["ok"] and events[-1]["chapters"] == 2
    request["protocolVersion"] = "bad"
    result = subprocess.run([sys.executable, "-m", "document_exporter"], input=json.dumps(request).encode(), capture_output=True, timeout=30, cwd=Path(__file__).resolve().parents[1])
    assert result.returncode == 1
    assert not json.loads(result.stdout)["ok"]


def test_invalid_content_never_reports_success(tmp_path: Path) -> None:
    source = book()
    source["volumes"][0]["chapters"][0]["paragraphs"] = ["bad\x00text"]
    with pytest.raises(ValueError, match="第 1 卷 江湖 / 第 1 章 重名"):
        export_document(source, "docx", tmp_path / "bad.docx")
    source["volumes"][0]["chapters"][0]["paragraphs"] = [42]
    with pytest.raises(ValueError, match="段落无效"):
        export_document(source, "txt", tmp_path / "bad.txt")


@pytest.mark.skipif(not os.environ.get("NOVEL_EXPORTER_BINARY"), reason="Packaged worker path was not supplied")
def test_packaged_worker_without_system_python(tmp_path: Path) -> None:
    (tmp_path / "snapshot.json").write_text(json.dumps(book(), ensure_ascii=False), encoding="utf-8")
    environment = {key: value for key, value in os.environ.items() if not key.startswith("PYTHON") and key not in ("VIRTUAL_ENV", "PATH")}
    environment["PATH"] = str(Path(os.environ["SystemRoot"]) / "System32") if os.name == "nt" else "/usr/bin:/bin"
    for format in ("txt", "docx"):
        request = {"protocolVersion": PROTOCOL, "jobId": "packaged", "directory": str(tmp_path), "format": format}
        result = subprocess.run([os.environ["NOVEL_EXPORTER_BINARY"]], input=json.dumps(request).encode(), capture_output=True, timeout=90, env=environment)
        assert result.returncode == 0, result.stderr
        response = json.loads(result.stdout.decode().splitlines()[-1])
        assert response["ok"] and response["chapters"] == 2
        output = tmp_path / f"output.{format}"
        if format == "txt":
            assert "　　正文😀  " in output.read_text(encoding="utf-8-sig")
        else:
            assert "　　正文😀  " in [p.text for p in Document(output).paragraphs]


def test_large_book_no_truncation(tmp_path: Path) -> None:
    source = deepcopy(book())
    source["volumes"] = [
        {"id": f"v{v}", "title": f"第 {v + 1} 卷", "chapters": [
            {"id": f"c{v}-{c}", "title": f"第 {v * 100 + c + 1} 章", "paragraphs": ["正文😀" * 400, f"章末-{v}-{c}"]}
            for c in range(100)
        ]} for v in range(10)
    ]
    for format in ("txt", "docx"):
        start = perf_counter()
        output = tmp_path / f"large.{format}"
        result = export_document(source, format, output)
        assert result["chapters"] == 1000
        text = output.read_text(encoding="utf-8-sig") if format == "txt" else "\n".join(p.text for p in Document(output).paragraphs)
        assert text.count("正文😀") == 400_000
        assert "章末-9-99" in text
        print(f"{format}: {perf_counter() - start:.2f}s, {output.stat().st_size} bytes")
    if os.name == "nt":
        import ctypes
        from ctypes import wintypes

        class Counters(ctypes.Structure):
            _fields_ = [("cb", wintypes.DWORD), ("faults", wintypes.DWORD)] + [(key, ctypes.c_size_t) for key in (
                "peakWorkingSet", "workingSet", "peakPaged", "paged", "peakNonPaged", "nonPaged", "pagefile", "peakPagefile")]

        counters = Counters()
        counters.cb = ctypes.sizeof(counters)
        ctypes.windll.kernel32.GetCurrentProcess.restype = wintypes.HANDLE
        query = ctypes.windll.psapi.GetProcessMemoryInfo
        query.argtypes = [wintypes.HANDLE, ctypes.POINTER(Counters), wintypes.DWORD]
        assert query(ctypes.windll.kernel32.GetCurrentProcess(), ctypes.byref(counters), counters.cb)
        print(f"pytest process peak working set: {counters.peakWorkingSet / 1024 / 1024:.1f} MiB")
