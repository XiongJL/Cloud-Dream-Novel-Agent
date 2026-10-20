from __future__ import annotations

from pathlib import Path
from typing import Any, Callable


PROTOCOL = "document-exporter-v1"


def validate_snapshot(snapshot: Any) -> dict[str, Any]:
    if not isinstance(snapshot, dict):
        raise ValueError("作品快照无效")
    for key in ("title", "author", "description"):
        if not isinstance(snapshot.get(key), str):
            raise ValueError(f"作品信息无效：{key}")
    if not isinstance(snapshot.get("volumes"), list):
        raise ValueError("作品卷列表无效")
    for volume in snapshot["volumes"]:
        if not isinstance(volume, dict) or not isinstance(volume.get("title"), str) or not isinstance(volume.get("chapters"), list):
            raise ValueError("卷数据无效")
        for chapter in volume["chapters"]:
            if not isinstance(chapter, dict) or not isinstance(chapter.get("title"), str):
                raise ValueError("章节数据无效")
            if not isinstance(chapter.get("paragraphs"), list) or not all(isinstance(p, str) for p in chapter["paragraphs"]):
                raise ValueError("章节段落无效")
    return snapshot


def _metadata(snapshot: dict[str, Any]) -> list[str]:
    blocks = [snapshot["title"]]
    if snapshot["author"].strip():
        blocks.append(f"作者：{snapshot['author']}")
    if snapshot["description"].strip():
        blocks.extend(("作品简介", snapshot["description"].replace("\r\n", "\n").replace("\r", "\n")))
    return blocks


def _write_txt(snapshot: dict[str, Any], output: Path, progress: Callable[[int], None]) -> None:
    with output.open("w", encoding="utf-8-sig", newline="") as stream:
        first = True
        last = ""

        def block(text: str) -> None:
            nonlocal first, last
            if not first:
                stream.write("\r\n\r\n")
            last = text.replace("\r\n", "\n").replace("\r", "\n").replace("\n", "\r\n")
            stream.write(last)
            first = False

        for text in _metadata(snapshot):
            block(text)
        completed = 0
        for volume in snapshot["volumes"]:
            block(volume["title"])
            for chapter in volume["chapters"]:
                block(chapter["title"])
                if chapter["paragraphs"]:
                    block("\n".join(chapter["paragraphs"]))
                completed += 1
                progress(completed)
        if not last.endswith("\r\n"):
            stream.write("\r\n")


def _write_docx(snapshot: dict[str, Any], output: Path, progress: Callable[[int], None]) -> None:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    doc.core_properties.title = snapshot["title"]
    doc.core_properties.author = snapshot["author"]
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.54)
    section.left_margin = section.right_margin = Cm(2.54)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    for name in ("Heading 1", "Heading 2"):
        doc.styles[name].paragraph_format.keep_with_next = True
    for name in ("TOC 1", "TOC 2"):
        if name not in doc.styles:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = normal
    doc.styles["TOC 2"].paragraph_format.left_indent = Cm(0.6)

    doc.add_paragraph(snapshot["title"], "Title")
    if snapshot["author"].strip():
        doc.add_paragraph(f"作者：{snapshot['author']}")
    if snapshot["description"].strip():
        doc.add_paragraph("作品简介", "Subtitle")
        for paragraph in snapshot["description"].replace("\r\n", "\n").replace("\r", "\n").split("\n"):
            doc.add_paragraph(paragraph)

    entries: list[tuple[int, str, str, int]] = []
    volume_bookmarks: list[tuple[dict[str, Any], tuple[str, int], list[tuple[dict[str, Any], tuple[str, int]]]]] = []
    for volume in snapshot["volumes"]:
        index = len(entries)
        bookmark = (f"ne_{index}", index)
        entries.append((1, volume["title"], *bookmark))
        chapters = []
        for chapter in volume["chapters"]:
            index = len(entries)
            chapter_bookmark = (f"ne_{index}", index)
            entries.append((2, chapter["title"], *chapter_bookmark))
            chapters.append((chapter, chapter_bookmark))
        volume_bookmarks.append((volume, bookmark, chapters))

    if entries:
        doc.add_paragraph("目录", "Subtitle")
        field = doc.add_paragraph()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        begin.set(qn("w:dirty"), "true")
        field.add_run()._r.append(begin)
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = ' TOC \\o "1-2" \\h \\z \\u '
        field.add_run()._r.append(instruction)
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        field.add_run()._r.append(separate)
        for level, title, bookmark_name, _ in entries:
            paragraph = doc.add_paragraph(style=f"TOC {level}")
            link = OxmlElement("w:hyperlink")
            link.set(qn("w:anchor"), bookmark_name)
            link.set(qn("w:history"), "1")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.set(qn("xml:space"), "preserve")
            text.text = title
            run.append(text)
            link.append(run)
            paragraph._p.append(link)
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        paragraph.add_run()._r.append(end)
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        doc.settings.element.append(update)

    def heading(title: str, level: int, bookmark: tuple[str, int]) -> None:
        paragraph = doc.add_heading(title, level=level)
        if level == 1:
            paragraph.paragraph_format.page_break_before = True
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(bookmark[1]))
        start.set(qn("w:name"), bookmark[0])
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(bookmark[1]))
        paragraph._p.insert(1, start)
        paragraph._p.append(end)

    completed = 0
    for volume, bookmark, chapters in volume_bookmarks:
        heading(volume["title"], 1, bookmark)
        for chapter, chapter_bookmark in chapters:
            heading(chapter["title"], 2, chapter_bookmark)
            try:
                for text in chapter["paragraphs"]:
                    doc.add_paragraph(text)
            except ValueError as error:
                raise ValueError(f"{volume['title']} / {chapter['title']}：正文包含 DOCX 无法表示的字符") from error
            completed += 1
            progress(completed)
    doc.save(output)


def export_document(snapshot: Any, format: str, output: Path, progress: Callable[[int], None] = lambda _: None) -> dict[str, Any]:
    data = validate_snapshot(snapshot)
    if format not in ("txt", "docx") or output.suffix != f".{format}":
        raise ValueError("不支持的导出格式")
    if format == "txt":
        _write_txt(data, output, progress)
    else:
        _write_docx(data, output, progress)
    return {
        "chapters": sum(len(v["chapters"]) for v in data["volumes"]),
        "volumes": len(data["volumes"]),
        "empty": not any(p.strip() for v in data["volumes"] for c in v["chapters"] for p in c["paragraphs"]),
    }
